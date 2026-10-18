# Extract text (body, tables, headers, footers) from tracked .docx/.pdf and run the dictionary. Out: out/doc_hits.json, out/doc_info.json
import json, os, re, subprocess, sys, zipfile
sys.path.insert(0, os.path.dirname(__file__))
from terms import COMPILED
import docx
from common import REPO, WORK as S, tracked_files
files = [f for f in tracked_files() if f.lower().endswith((".docx", ".pdf"))]
os.makedirs(os.path.join(S, "docs"), exist_ok=True)
hits, info = [], {}
for f in files:
    p = os.path.join(REPO, f)
    parts = []
    if f.endswith(".docx"):
        d = docx.Document(p)
        for s in d.sections:
            for hp in (s.header, s.footer, s.first_page_header, s.first_page_footer):
                parts += ["[header/footer] " + x.text for x in hp.paragraphs if x.text.strip()]
        parts += [x.text for x in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text for c in r.cells))
        z = zipfile.ZipFile(p)
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        hdr_rels = [n for n in z.namelist() if re.match(r"word/_rels/(header|footer)\d*\.xml\.rels", n)]
        hdr_imgs = sorted({m for n in hdr_rels for m in re.findall(r'Target="media/([^"]+)"', z.read(n).decode())})
        info[f] = {"media": len(media), "header_images": hdr_imgs}
    else:
        txt = subprocess.check_output(["pdftotext", "-layout", p, "-"]).decode("utf-8", "replace")
        parts = txt.split("\n")
        info[f] = {"pages": int(re.search(r"Pages:\s+(\d+)", subprocess.check_output(["pdfinfo", p]).decode()).group(1))}
    text = "\n".join(parts)
    open(os.path.join(S, "docs", f.replace("/", "_") + ".txt"), "w").write(text)
    for n, line in enumerate(text.split("\n"), 1):
        for term, rx, cat, origin in COMPILED:
            for m in rx.finditer(line):
                hits.append({"term": term, "file": f, "line": n, "match": m.group(0), "text": line.strip()[:200]})
json.dump(hits, open(os.path.join(S, "out", "doc_hits.json"), "w"), ensure_ascii=False, indent=0)
json.dump(info, open(os.path.join(S, "out", "doc_info.json"), "w"), ensure_ascii=False, indent=1)
import collections
print(len(files), "docs;", len(hits), "hits")
print(collections.Counter(h["term"] for h in hits).most_common(30))
